"""生成 NOVEL-AGENT 全景介绍 16:9 高清商务 PPT (PowerPoint .pptx)。

视觉规范：
- 16:9 宽屏布局 (13.333 x 7.5 英寸)
- 极客科技暗黑美学 (底色 #080C0A，翡翠绿强调色 #10B981 / #34D399，暗灰对比色 #1E293B)
- 每页附带演讲者备注 (Speaker Notes)，支持 PowerPoint 演讲者视图
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# 颜色定义
BG_COLOR = RGBColor(8, 12, 10)         # 深沉暗黑
CARD_BG = RGBColor(15, 23, 20)         # 卡片背景
CARD_BORDER = RGBColor(20, 60, 45)     # 卡片边框微光
EMERALD_BRIGHT = RGBColor(52, 211, 153) # 亮翡翠绿
EMERALD_MAIN = RGBColor(16, 185, 129)   # 主翡翠绿
EMERALD_DARK = RGBColor(6, 78, 59)     # 暗翡翠绿
TEXT_WHITE = RGBColor(248, 250, 252)   # 标题白
TEXT_SLATE = RGBColor(148, 163, 184)   # 副标题灰
TEXT_MUTED = RGBColor(100, 116, 139)   # 次要灰
RED_ALERT = RGBColor(239, 68, 68)      # 警示红
RED_BG = RGBColor(40, 18, 18)          # 警示卡片底色
BLUE_ACCENT = RGBColor(56, 189, 248)   # 科技蓝
PURPLE_ACCENT = RGBColor(168, 85, 247) # 紫色

def set_slide_background(slide):
    """设置幻灯片深暗背景。"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, tag_text: str, title_text: str, subtitle_text: str = ""):
    """添加统一样式的顶部标题栏。"""
    # 顶部标签
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
    tf = tag_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = tag_text.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD_BRIGHT
    p.font.name = "Consolas"

    # 大标题
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.font.name = "Microsoft YaHei"

    # 副标题
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.7), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = TEXT_SLATE
        p_sub.font.name = "Microsoft YaHei"

def add_card(slide, left, top, width, height, title, body_lines, border_color=CARD_BORDER, bg_color=CARD_BG, title_color=EMERALD_BRIGHT):
    """绘制高质感圆角信息卡片。"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.2)

    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.22)
    tf.margin_bottom = Inches(0.2)

    # 卡片标题
    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.bold = True
    p_title.font.size = Pt(15)
    p_title.font.color.rgb = title_color
    p_title.font.name = "Microsoft YaHei"
    p_title.space_after = Pt(8)

    # 正文内容
    for idx, line in enumerate(body_lines):
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_SLATE
        p.font.name = "Microsoft YaHei"
        p.space_after = Pt(4)

def set_notes(slide, notes_text: str):
    """注入演讲者备忘录（逐字稿）。"""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = f"【演讲者逐字稿 / 提示】：\n{notes_text.strip()}"

def build_presentation(output_path: str):
    prs = Presentation()
    # 设置 16:9 比例 (13.333 x 7.5 英寸)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # 空白布局

    # =========================================================================
    # SLIDE 1: 封面
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # 徽章胶囊
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.3), Inches(1.5), Inches(4.7), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = EMERALD_DARK
    badge.line.color.rgb = EMERALD_MAIN
    badge.line.width = Pt(1)
    b_tf = badge.text_frame
    b_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = b_tf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = "确定性状态内核 × 六维智能审稿 × 极客全沉浸 IDE"
    bp.font.size = Pt(11)
    bp.font.bold = True
    bp.font.color.rgb = TEXT_WHITE
    bp.font.name = "Consolas"

    # 主标题
    title_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(1.4))
    tf1 = title_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "NOVEL-AGENT"
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = EMERALD_BRIGHT
    p1.font.name = "Consolas"

    # 副标题
    sub_box = s1.shapes.add_textbox(Inches(1.5), Inches(3.5), Inches(10.333), Inches(0.8))
    tf1_sub = sub_box.text_frame
    p1_sub = tf1_sub.paragraphs[0]
    p1_sub.alignment = PP_ALIGN.CENTER
    p1_sub.text = "长篇小说写作状态引擎 · 解决 AI 长篇创作断代遗忘的底层基础设施"
    p1_sub.font.size = Pt(18)
    p1_sub.font.color.rgb = TEXT_WHITE
    p1_sub.font.name = "Microsoft YaHei"

    p1_slogan = tf1_sub.add_paragraph()
    p1_slogan.alignment = PP_ALIGN.CENTER
    p1_slogan.text = "“用严谨的工程与状态机，护航狂野的文学创造力”"
    p1_slogan.font.size = Pt(14)
    p1_slogan.font.color.rgb = EMERALD_MAIN
    p1_slogan.space_before = Pt(8)

    # 底部 4 个关键特性亮点
    c_w = Inches(2.7)
    c_h = Inches(1.6)
    c_top = Inches(4.8)
    add_card(s1, Inches(0.8), c_top, c_w, c_h, "01 / 状态持久", ["大五人格按章演变", "角色关系动态网", "突变率阈值安全护栏"])
    add_card(s1, Inches(3.8), c_top, c_w, c_h, "02 / 伏笔闭环", ["双向精准锚定台账", "开放/收束状态机", "拒绝烂尾与断头悬念"])
    add_card(s1, Inches(6.8), c_top, c_w, c_h, "03 / 六维审稿", ["修辞/张力/因果诊断", "三级修改梯度分流", "Diff 对比就地原子写回"])
    add_card(s1, Inches(9.8), c_top, c_w, c_h, "04 / 极客 IDE", ["打字机垂直平滑居中", "单文件 264KB 免安装", "双重视角自由拖拽切换"])

    set_notes(s1, "各位好！今天向大家汇报的项目是 NOVEL-AGENT。在当今大语言模型写作蔚然成风的时代，为什么市面上依然没有一部真正由 AI 稳定写出的优质长篇小说？因为所有大模型都有一个致命的阿喀琉斯之踵。我们正是为了解决长篇状态遗忘与人设崩塌而诞生的专业状态引擎。")

    # =========================================================================
    # SLIDE 2: 行业痛点
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "02 // 行业痛点与背景", "大语言模型长篇创作的“四大死穴”", "LLM 天生是极佳的短篇扩写者，但受限于注意力衰减与无状态机制，充当不了“小说家的大脑”。")

    card_w2 = Inches(5.6)
    card_h2 = Inches(2.2)
    add_card(s2, Inches(0.8), Inches(2.3), card_w2, card_h2, "❌ 01. 上下文断代遗忘 (Context Loss)", [
        "• 长篇小说跨度数十万至数百万字，远超大模型上下文窗口承载极限。",
        "• 写到第 20 章早已忘光第 3 章埋设的暗线；强行塞入全本大纲不仅产生高昂 Token 费用，还会因注意力分散导致关键指令被稀释忽略。"
    ], border_color=RGBColor(120, 30, 30), bg_color=RED_BG, title_color=RED_ALERT)

    add_card(s2, Inches(6.8), Inches(2.3), card_w2, card_h2, "❌ 02. 角色人设崩塌漂移 (OOC)", [
        "• 缺乏结构化的性格与心理状态基准线追踪机制。",
        "• 主角前一章刚毅果断，后一章因外部事件突变为软弱圣母；配角动机飘忽不定，人物台词严重脱离既有经历背景。"
    ], border_color=RGBColor(120, 30, 30), bg_color=RED_BG, title_color=RED_ALERT)

    add_card(s2, Inches(0.8), Inches(4.7), card_w2, card_h2, "❌ 03. 伏笔悬空与因果悖论", [
        "• 无法精准追踪线索生命周期（埋设 -> 推进 -> 揭秘收束）。",
        "• 前期铺垫的宝物或伏笔成为无人认领的悬空断头线，前后因果逻辑互斥，最终只能依赖粗暴的“机械降神”仓促收场。"
    ], border_color=RGBColor(120, 30, 30), bg_color=RED_BG, title_color=RED_ALERT)

    add_card(s2, Inches(6.8), Inches(4.7), card_w2, card_h2, "❌ 04. 全知 Prompt 幻觉与假成功", [
        "• 市面传统写作工具往往将状态管理甩给单次 Prompt 盲目生成。",
        "• 许多工具将模型调用失败静默吞成假成功；缺乏工业级防御隔离，模型胡编乱造越积越多，最终导致整本小说逻辑彻底崩塌。"
    ], border_color=RGBColor(120, 30, 30), bg_color=RED_BG, title_color=RED_ALERT)

    set_notes(s2, "传统 AI 写作工具的做法是：拼命把几万字的大纲和背景塞进 Prompt 问模型：‘请写第 15 章’。结果往往是灾难性的：模型不仅吞噬巨额费用，还会因注意力机制衰减而顾此失彼。第 15 章里的主角性格突然大变，第 3 章精心设计的伏笔也永远无人回收。长篇写作最需要的不是堆砌 Token，而是精密的因果秩序。")

    # =========================================================================
    # SLIDE 3: 核心设计哲学
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "03 // 破局思考与设计哲学", "确定性领域内核与生成能力分离", "“大语言模型负责妙笔生花的文学扩写，确定性状态机负责因果台账的铁律守护。”")

    c_w3 = Inches(3.7)
    c_h3 = Inches(4.6)
    add_card(s3, Inches(0.8), Inches(2.2), c_w3, c_h3, "① 领域内核零 LLM 污染", [
        "【架构免疫系统】",
        "• src/novel/ 下四大核心引擎（角色/剧情/项目/章节组装）绝不导入任何大模型依赖。",
        "• 纯确定性算法推导：大五人格演变、关系矩阵向量、伏笔台账均为确定性数据结构。",
        "• 保证无论上层 LLM 如何产生幻觉，底层作品状态库绝对保真、不被污染。"
    ], border_color=CARD_BORDER, bg_color=CARD_BG, title_color=EMERALD_BRIGHT)

    add_card(s3, Inches(4.8), Inches(2.2), c_w3, c_h3, "② 唯一出口与模型无关", [
        "【模型配置即数据】",
        "• 全仓所有 LLM 调用统一收敛至 src/llm/ 唯一出口，业务代码零直接 import。",
        "• FastPath 策略：Flash 极速伴写档（大纲初稿）+ Pro 深度推理档（深度写作与六维审稿）。",
        "• 运行时动态注入：支持热切换 DeepSeek、Claude、OpenAI 等任意兼容端点。"
    ], border_color=CARD_BORDER, bg_color=CARD_BG, title_color=BLUE_ACCENT)

    add_card(s3, Inches(8.8), Inches(2.2), c_w3, c_h3, "③ 唯一操作契约与护栏", [
        "【彻底根绝“假成功”】",
        "• 所有业务逻辑由 NovelService 统一门面驱动，Web / TUI / CLI 均是轻量适配层。",
        "• 严谨失败语义：以 Result(ok=False) 显式表达，杜绝静默假成功。",
        "• 数据落盘护栏：正文写回前强制自动备份带时间戳的 .bak 镜像，支持无损回滚。"
    ], border_color=CARD_BORDER, bg_color=CARD_BG, title_color=PURPLE_ACCENT)

    set_notes(s3, "我们的破局思考非常清晰：大模型天然会写优美的段落，但大模型根本不配、也不应该作为‘小说数据库’！真正的状态——谁和谁好感度是多少、有哪些未解悬念、章节的时间先后顺序——必须交给 100% 可靠的确定性引擎。让确定性代码负责记账与推导，让大模型只负责最后一步的文学扩写与修辞。")

    # =========================================================================
    # SLIDE 4: 核心技术创新一
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "04 // 核心技术突破一", "大五人格连续演变与伏笔闭环台账", "用量化心理学与结构化生命周期，让角色拥有连续灵魂，让剧情杜绝断头烂尾。")

    c_w4 = Inches(5.6)
    c_h4 = Inches(4.6)
    add_card(s4, Inches(0.8), Inches(2.2), c_w4, c_h4, "● 大五人格 (OCEAN) 连续演变模型", [
        "【科学量化人物心理状态】",
        "• 建立开放性、尽责性、外倾性、宜人性、神经质五维数值（0-10）。",
        "• 逐章事件触发演化：记录具体受创、转折与成长事件，动态更新特质数值。",
        "• 突变率阈值护栏（<15%）：单章特质变动过大时自动预警并平滑夹紧，拦截人设断崖崩溃。",
        "• 定向动态关系网：量化人物好感与敌对向量，精确记录起始章节与最近变动原因。",
        "• 智能写作切片：仅将当前章节登场角色的心理账单精准注入 Prompt，实现极小 Token 成本的高质量角色还原。"
    ])

    add_card(s4, Inches(6.8), Inches(2.2), c_w4, c_h4, "★ 伏笔双向追踪台账 (Foreshadowing Ledger)", [
        "【长线叙事防烂尾引擎】",
        "• 严格三阶生命周期状态机：[开放 Open] -> [推进 Progress] -> [收束 Resolved]。",
        "• 全维度元数据登记：埋设章节、预期揭秘章节、关联关键角色、线索重要等级。",
        "• 双向闭环检索：在企划界面点击伏笔秒级定位对应章节正文；反向查看章节时高亮其背负的收束使命。",
        "• 写作护栏校验：章节生成前自动扫描活跃伏笔，提示模型在此处铺垫或引爆，杜绝草蛇灰线被长久遗忘。"
    ])

    set_notes(s4, "在这个系统里，角色不是一张静态的说明书，而是一个活生生的状态机。第 12 章主角遭遇背叛，他的宜人性和信任值就地扣减并记录到事件流中；到了第 30 章写作时，上下文引擎会自动提取这笔心理账单，提示模型：‘该角色处于创伤戒备期，面对好意时绝不轻易相信’。伏笔台账更是确保每一把挂在墙上的猎枪都能在合情合理的位置打响。")

    # =========================================================================
    # SLIDE 5: 核心技术创新二
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "05 // 核心技术突破二", "六维智能审稿与就地原子改写", "拒绝空洞套话，提供工业级“精准锚点搜索、前后对比 Diff 与一键改写”闭环流水线。")

    # 6个维度网格
    grid_w = Inches(3.7)
    grid_h = Inches(1.8)
    add_card(s5, Inches(0.8), Inches(2.2), grid_w, grid_h, "1. 文笔修辞与感官描写", [
        "• 诊断白描冗余、成语堆砌误用",
        "• 检视视听触嗅五感层次丰富度"
    ])
    add_card(s5, Inches(4.8), Inches(2.2), grid_w, grid_h, "2. 剧情张力与场景动力", [
        "• 检验场景价值转换（正转负/负转正）",
        "• 拦截悬念衰减与无效对话灌水"
    ])
    add_card(s5, Inches(8.8), Inches(2.2), grid_w, grid_h, "3. 伏笔闭环与因果自洽", [
        "• 跨章节核对主线时间线单调性",
        "• 杜绝逻辑悖论与生硬机械降神"
    ])
    add_card(s5, Inches(0.8), Inches(4.2), grid_w, grid_h, "4. 角色人设与台词质感", [
        "• 比对角色大五人格基准线",
        "• 拦截非预期性格突变与脸谱化台词"
    ])
    add_card(s5, Inches(4.8), Inches(4.2), grid_w, grid_h, "5. 世界观自洽与战力平衡", [
        "• 校验修炼阶梯、科技法则是否逾矩",
        "• 监控战力系统膨胀与等级崩坏"
    ])
    add_card(s5, Inches(8.8), Inches(4.2), grid_w, grid_h, "6. 叙事节奏与详略剪裁", [
        "• 章节波峰波谷叙事波形分析",
        "• 识别平铺直叙与节奏拖沓淤塞"
    ])

    # 底部横条：三级梯队与安全保护
    alert_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.65))
    alert_box.fill.solid()
    alert_box.fill.fore_color.rgb = CARD_BG
    alert_box.line.color.rgb = EMERALD_MAIN
    atf = alert_box.text_frame
    atf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ap = atf.paragraphs[0]
    ap.text = "⚡ 分级治理闭环：Level 1 机械锚点一键替换 · Level 2 原地 Diff 采纳 · 覆盖前自动生成 project.json.bak 镜像"
    ap.font.size = Pt(11)
    ap.font.bold = True
    ap.font.color.rgb = EMERALD_BRIGHT

    set_notes(s5, "许多 AI 审稿工具只会给你一段毫无意义的‘文笔优美、建议增强心理活动’套话。NOVEL-AGENT 的六维审稿是一个真正的工业级质检流水线。它会并发从修辞、节奏、因果律、人设等 6 个维度切片诊断，并直接给出具体的修改前与修改后文本（Diff）。作者在界面上只需点一下‘采纳’，系统就会在自动备份原稿的同时，将修改安全注入作品。")

    # =========================================================================
    # SLIDE 6: 核心产品呈现
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "06 // 核心产品呈现", "Novel Studio 现代沉浸式写作 IDE", "复古单色 CRT 终端美学 × 极致舒适的现代文学排版工作台。")

    c_w6 = Inches(5.6)
    c_h6 = Inches(3.6)
    add_card(s6, Inches(0.8), Inches(2.2), c_w6, c_h6, "【企划构思视角 (Planning Perspective)】", [
        "• 默认首页直达：首屏全局纵览书名、题材、全书篇幅及大纲骨架。",
        "• 分卷与章节树状目录：已写/待写章节、字数与剧情概要一览无余。",
        "• 出场角色时间切片：选中章节自动过滤该章登场人物与其心理雷达图。",
        "• 伏笔台账双向穿梭：点击伏笔秒级定位对应章节与高潮收束点。"
    ])

    add_card(s6, Inches(6.8), Inches(2.2), c_w6, c_h6, "【沉浸写作视角 (Writing Perspective)】", [
        "• 打字机垂直平滑居中：光标始终锁定屏幕正中央，告别低头视线疲劳。",
        "• 左右工作区自由伸缩：鼠标任意拖拽分割中缝，自由定制双栏比例。",
        "• 专业文学排版定制：字号（15~21px）、行距、专业宋体/楷体无缝切换。",
        "• 视效主题自由换肤：黑客翠绿、琥珀金黄、赛博青蓝、极简黑白 4 套配色。"
    ])

    # 底部绿色高亮条
    bot_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.85))
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = CARD_BG
    bot_card.line.color.rgb = EMERALD_BRIGHT
    btf = bot_card.text_frame
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = btf.paragraphs[0]
    bp.text = "🚀 单文件独立便携版已发布 (264 KB HTML)：双击即可在任意浏览器离线写作，无需安装 Python 环境，草稿全在本地持久化！"
    bp.font.size = Pt(12)
    bp.font.bold = True
    bp.font.color.rgb = TEXT_WHITE

    set_notes(s6, "不仅底层架构过硬，在交互体验上，我们精心打磨了 Novel Studio。这是一座带有复古单色 CRT 终端美学、同时又具备极度现代文学排版的工作台。它拥有作者梦寐以求的打字机居中模式、自由伸缩的分屏，以及单文件开箱即用的便携体验。无论是在高性能工作站还是轻薄笔记本上，都能提供丝滑的创作享受。")

    # =========================================================================
    # SLIDE 7: 三位一体交付形态
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "07 // 三位一体交付矩阵", "全场景创作与开发者生态适配", "同一套确定性领域引擎，无缝驱动三种不同形态的交互界面。")

    c_w7 = Inches(3.7)
    c_h7 = Inches(3.6)
    add_card(s7, Inches(0.8), Inches(2.2), c_w7, c_h7, "Web 写作 IDE 版", [
        "【命令：uv run na-studio】",
        "• 适合日常主力创作与精细化排版。",
        "• 实时图形化大纲与五维雷达图联动。",
        "• 支持导出整本首行缩进 TXT、全案 Markdown 以及全套 JSON 备份。",
        "• 自带 RESTful API 服务端，支持局域网或云端私有化部署。"
    ], border_color=CARD_BORDER, bg_color=CARD_BG, title_color=BLUE_ACCENT)

    add_card(s7, Inches(4.8), Inches(2.2), c_w7, c_h7, "TUI 终端交互版", [
        "【命令：uv run na-tui】",
        "• 基于 Textual 框架打造的现代终端应用。",
        "• 纯键盘全快捷键驱动，极速响应。",
        "• 适合远程 Linux / SSH 无头服务器写作。",
        "• 审稿面板按 Enter 原地对比 Diff 并一键确认采纳写回正文。"
    ], border_color=CARD_BORDER, bg_color=CARD_BG, title_color=EMERALD_BRIGHT)

    add_card(s7, Inches(8.8), Inches(2.2), c_w7, c_h7, "CLI 自动化命令版", [
        "【命令：uv run na novel ...】",
        "• 适合自动化流水线与批处理脚本。",
        "• na novel write：单章指定字数与档位生成。",
        "• na novel auto：从首个未写章节连续续写。",
        "• na novel review：命令行非交互式审稿。",
        "• na kb search：写作语料与设定检索。"
    ], border_color=CARD_BORDER, bg_color=CARD_BG, title_color=PURPLE_ACCENT)

    # 底部未来生态演进
    mcp_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.85))
    mcp_card.fill.solid()
    mcp_card.fill.fore_color.rgb = CARD_BG
    mcp_card.line.color.rgb = PURPLE_ACCENT
    mtf = mcp_card.text_frame
    mtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    mp = mtf.paragraphs[0]
    mp.text = "🔮 未来生态演进 (MCP 服务器)：通过 Model Context Protocol 标准协议，直接成为 Claude / Cursor 等桌面宿主 AI 的长篇小说外脑！"
    mp.font.size = Pt(12)
    mp.font.bold = True
    mp.font.color.rgb = TEXT_WHITE

    set_notes(s7, "针对不同的创作者人群与使用场景，我们提供了完整的三位一体交互能力：创作者在日常使用时可以用 Web 享受排版与视效；在轻量终端或 VPS 上可以用 TUI 全键盘操作；运维与自动化流程中可以用 CLI 一行命令批量续写或审稿；而未来通过 MCP，它将成为 Claude 等顶级宿主的专属小说记忆芯片。")

    # =========================================================================
    # SLIDE 8: 工业级工程质量保障
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "08 // 工业级工程质量保障", "全量离线回归与严苛 CI 质量门禁", "以对待数据库与金融系统的工程严谨度，保障长篇小说的每一行文字资产。")

    c_w8 = Inches(5.6)
    c_h8 = Inches(3.6)
    add_card(s8, Inches(0.8), Inches(2.2), c_w8, c_h8, "● 360+ 离线测试套件 (Zero Cost)", [
        "【361 passed, 0 failed in 23s】",
        "• 覆盖角色状态演变、剧情约束、审稿算法、服务层及 API 路由。",
        "• 硬规则防线：测试统一使用 TestModel 测试桩，绝不发起真实外网调用，离线运行零计费消耗。",
        "• 契约锁定：对 characters.json / project.json 历史读写做基线锁定，严防破坏向下兼容。"
    ])

    add_card(s8, Inches(6.8), Inches(2.2), c_w8, c_h8, "● GitHub Actions CI 云端三重门禁", [
        "【全天候自动化流水线保障】",
        "• Ruff 静态代码检查：统一 Python 3.10 语法标准，严格执行行宽 100 规范。",
        "• Mypy 强类型检查：对核心可靠性模块执行类型检查，杜绝动态类型隐患。",
        "• 全回归自动化测试：任何 Pull Request 与 Commit 必须 100% 绿灯方可合并。"
    ])

    # 底部数据安全
    sec_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.85))
    sec_card.fill.solid()
    sec_card.fill.fore_color.rgb = CARD_BG
    sec_card.line.color.rgb = EMERALD_MAIN
    stf = sec_card.text_frame
    stf.vertical_anchor = MSO_ANCHOR.MIDDLE
    sp = stf.paragraphs[0]
    sp.text = "🛡️ 数据绝对安全：采用 os.replace 原子写盘机制，写入前强制生成带毫秒时间戳的 .bak 镜像，断电与异常死机永不丢稿！"
    sp.font.size = Pt(12)
    sp.font.bold = True
    sp.font.color.rgb = TEXT_WHITE

    set_notes(s8, "做小说引擎，最可怕的事情莫过于‘写着写着作者的稿子没了’。在工程层面上，我们有着近乎偏执的安全追求。全项目 361 个单元测试全部实现零网络、零费用的纯离线断言；每一次存盘都采用原子文件替换与自动备份双重保险；所有代码经过 GitHub CI 的 Ruff 和 Mypy 严苛审查，确保绝对的工业级稳定性。")

    # =========================================================================
    # SLIDE 9: 路线图与演进规划
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "09 // 路线图与演进规划", "从单机写作工具走向文学创作操作系统", "分阶段演进路径清晰，持续开拓下一代人机协同叙事范式。")

    rm_w = Inches(11.7)
    add_card(s9, Inches(0.8), Inches(2.2), rm_w, Inches(1.2), "【已完成】Phase 1-2：确定性领域内核与 Novel Studio 体验革命", [
        "• 角色大五人格演化机、动态关系网、伏笔双向追踪台账落盘验证。",
        "• 六维智能审稿流水线与 Diff 对比原子写回就绪；Novel Studio 现代 Web IDE 正式发布 (v0.1.0)。"
    ], border_color=EMERALD_MAIN, title_color=EMERALD_BRIGHT)

    add_card(s9, Inches(0.8), Inches(3.6), rm_w, Inches(1.2), "【进行中】Phase 3：标准化 MCP 服务器与宿主生态打通", [
        "• 落地标准 na-mcp 服务端（支持 Stdio 与 Streamable HTTP 双通道）。",
        "• 让外部 AI 宿主（Claude Desktop / Cursor）一键挂载小说引擎，把角色与伏笔账单作为外部状态调阅。"
    ], border_color=BLUE_ACCENT, title_color=BLUE_ACCENT)

    add_card(s9, Inches(0.8), Inches(5.0), rm_w, Inches(1.2), "【规划中】Phase 4：多智能体协作集群 (Agent Swarm) 与出版级排版输出", [
        "• 引入“黄金开篇教练”、“世界观巡检员”等专用专家 Agent 集群并行协作。",
        "• 研发出版级排版引擎，支持一键导出带精美扉页与目录的出版级 EPUB / PDF。"
    ], border_color=PURPLE_ACCENT, title_color=PURPLE_ACCENT)

    set_notes(s9, "目前，我们已经顺利完成了基础状态引擎的确定性构建、六维审稿闭环以及 Web IDE 的正式发布（v0.1.0）。接下来，我们将把系统全面打包为标准的 MCP 服务，让任何支持 MCP 的宿主 AI 都能立刻接入这颗强大的小说记忆引擎，实现真正由 AI 辅助编写的百万字鸿篇巨制。")

    # =========================================================================
    # SLIDE 10: 结语与 Q&A
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)

    # 结语徽章
    badge10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.3), Inches(1.5), Inches(4.7), Inches(0.45))
    badge10.fill.solid()
    badge10.fill.fore_color.rgb = EMERALD_DARK
    badge10.line.color.rgb = EMERALD_MAIN
    badge10.line.width = Pt(1)
    b10_tf = badge10.text_frame
    b10_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    b10_p = b10_tf.paragraphs[0]
    b10_p.alignment = PP_ALIGN.CENTER
    b10_p.text = "让长篇不再断代 · 让灵感触手可及"
    b10_p.font.size = Pt(11)
    b10_p.font.bold = True
    b10_p.font.color.rgb = TEXT_WHITE
    b10_p.font.name = "Consolas"

    # 主结语
    t10_box = s10.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(1.2))
    tf10 = t10_box.text_frame
    p10 = tf10.paragraphs[0]
    p10.alignment = PP_ALIGN.CENTER
    p10.text = "让技术回归秩序，让创作绽放光芒"
    p10.font.size = Pt(40)
    p10.font.bold = True
    p10.font.color.rgb = TEXT_WHITE
    p10.font.name = "Microsoft YaHei"

    sub10_p = tf10.add_paragraph()
    sub10_p.alignment = PP_ALIGN.CENTER
    sub10_p.text = "“大语言模型赋予我们文学扩写的羽翼，而确定性状态引擎赋予我们远航百里的航标与舵轮。”"
    sub10_p.font.size = Pt(16)
    sub10_p.font.color.rgb = EMERALD_BRIGHT
    sub10_p.font.name = "Microsoft YaHei"
    sub10_p.space_before = Pt(12)

    # 开源与下载信息卡片
    info_card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.6), Inches(4.0), Inches(6.1), Inches(1.8))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = CARD_BG
    info_card.line.color.rgb = CARD_BORDER
    itf = info_card.text_frame
    itf.margin_top = Inches(0.25)
    itf.margin_left = Inches(0.3)
    
    ip1 = itf.paragraphs[0]
    ip1.text = "📦 开源主仓库：https://github.com/wsljjjzn/NOVEL-AGENT"
    ip1.font.size = Pt(13)
    ip1.font.color.rgb = TEXT_WHITE
    ip1.font.name = "Consolas"
    
    ip2 = itf.add_paragraph()
    ip2.text = "🚀 Release 发布：v0.1.0 绿色单文件独立免安装版已就绪"
    ip2.font.size = Pt(13)
    ip2.font.color.rgb = EMERALD_BRIGHT
    ip2.font.name = "Microsoft YaHei"
    ip2.space_before = Pt(8)

    ip3 = itf.add_paragraph()
    ip3.text = "⚡ 快速体验：uv sync --extra web && uv run na-studio"
    ip3.font.size = Pt(13)
    ip3.font.color.rgb = TEXT_SLATE
    ip3.font.name = "Consolas"
    ip3.space_before = Pt(8)

    # QA
    qa_box = s10.shapes.add_textbox(Inches(2.0), Inches(6.1), Inches(9.333), Inches(0.6))
    qtf = qa_box.text_frame
    qp = qtf.paragraphs[0]
    qp.alignment = PP_ALIGN.CENTER
    qp.text = "THANK YOU // Q & A 问答交流环节"
    qp.font.size = Pt(15)
    qp.font.bold = True
    qp.font.color.rgb = EMERALD_BRIGHT
    qp.font.name = "Consolas"

    set_notes(s10, "最后，我想说的是，NOVEL-AGENT 的初心并不是用机器取代作家的创造力，而是把作家从繁杂的人物台账、伏笔记录和前后矛盾的记忆负担中彻底解放出来。让技术回归秩序，让创作绽放光芒。感谢大家！欢迎提问！")

    # 保存
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    print(f"[OK] 成功生成高清 16:9 演示文稿: {output_path}")

if __name__ == "__main__":
    out_file = "dist/NOVEL-AGENT-项目全景介绍.pptx"
    if len(sys.argv) > 1:
        out_file = sys.argv[1]
    build_presentation(out_file)
